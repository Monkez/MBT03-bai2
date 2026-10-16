"""Excel roster and pedestal assignments, independent of the scoring workers."""
import os
import tempfile
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from PyQt5 import uic
from PyQt5.QtCore import QEvent, QObject, Qt, pyqtSignal
from PyQt5.QtWidgets import QFileDialog, QMessageBox, QSizePolicy, QWidget

import config as cf
from gui.client_widget import classify_hits


RESULT_COLUMNS = ((1, "Bia số 6"), (0, "Bia số 10"), (2, "Bia số 7b"), (3, "Bia số 8c"))


@dataclass(frozen=True)
class RosterResult:
    hit_targets: frozenset[int]

    @property
    def rank(self):
        return classify_hits(self.hit_targets)


@dataclass(frozen=True)
class Shooter:
    number: int
    name: str
    unit: str
    birthday: str
    pedestal: int
    excel_row: int
    result: RosterResult | None = None


def _blank(value):
    return value is None or str(value).strip() == ""


def _normalized(value):
    text = unicodedata.normalize("NFD", str(value).strip().lower()).replace("đ", "d")
    return "".join(c for c in text if not unicodedata.combining(c) and not c.isspace())


def _validate_result_headers(sheet):
    for column, title in enumerate([title for _, title in RESULT_COLUMNS] + ["Xếp loại"], start=6):
        value = sheet.cell(1, column).value
        if not _blank(value) and _normalized(value) != _normalized(title):
            raise ValueError(f"Cột {chr(64 + column)} phải là '{title}', đang là '{value}'.")


def _read_result(values, row_number):
    if all(_blank(value) for value in values):
        return None
    hits = set()
    for value, (target, title) in zip(values[:4], RESULT_COLUMNS):
        if isinstance(value, bool) or str(value).strip() not in ("0", "1", "0.0", "1.0"):
            raise ValueError(f"Dòng {row_number}: {title} phải là 0 hoặc 1; "
                             "người chưa bắn để trống cả F–J.")
        if float(value) == 1:
            hits.add(target)
    result = RosterResult(frozenset(hits))
    if not _blank(values[4]) and _normalized(values[4]) != _normalized(result.rank):
        raise ValueError(f"Dòng {row_number}: xếp loại phải là '{result.rank}' theo kết quả bốn bia.")
    return result


def read_roster(path, pedestal_count):
    """Read identity A:E and optional target results F:J from the Bài 2 template."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    people = []
    numbers = set()
    try:
        _validate_result_headers(workbook.active)
        for row_number, row in enumerate(
                workbook.active.iter_rows(min_row=2, max_col=10, values_only=True), start=2):
            if all(_blank(value) for value in row):
                continue
            number, name, unit, birthday = row[:4]
            try:
                parsed = int(number)
                if isinstance(number, bool) or parsed <= 0 or float(number) != parsed:
                    raise ValueError()
            except (TypeError, ValueError, OverflowError):
                raise ValueError(f"Dòng {row_number}: STT phải là số nguyên dương.") from None
            if parsed in numbers:
                raise ValueError(f"Dòng {row_number}: STT {parsed} bị trùng.")
            if name is None or not str(name).strip():
                raise ValueError(f"Dòng {row_number}: thiếu họ tên.")
            birthday = (birthday.strftime("%d/%m/%Y") if isinstance(birthday, (date, datetime))
                        else str(birthday or "").strip())
            people.append(Shooter(parsed, str(name).strip(), str(unit or "").strip(),
                                  birthday, (parsed - 1) % pedestal_count + 1, row_number,
                                  _read_result(row[5:10], row_number)))
            numbers.add(parsed)
    finally:
        workbook.close()
    if not people:
        raise ValueError("File chưa có người bắn. Dòng 1 là tiêu đề; dữ liệu bắt đầu từ dòng 2.")
    return people


def save_roster_results(path, people, results):
    """Save a sibling workbook, keeping the input and previous output intact on failure."""
    source = Path(path)
    destination = source.with_name("Kết quả - " + source.name)
    workbook = load_workbook(source)
    temporary = None
    try:
        sheet = workbook.active
        _validate_result_headers(sheet)
        for person in people:
            row = person.excel_row
            number = sheet.cell(row, 1).value
            try:
                same_number = (not isinstance(number, bool) and int(number) == person.number
                               and float(number) == person.number)
            except (TypeError, ValueError, OverflowError):
                same_number = False
            if not same_number or str(sheet.cell(row, 2).value).strip() != person.name:
                raise ValueError("Danh sách trong file đã thay đổi. Nạp lại danh sách trước khi lưu kết quả.")
            result = results.get(person.number)
            for column, (target, _) in enumerate(RESULT_COLUMNS, start=6):
                sheet.cell(row, column).value = int(target in result.hit_targets) if result else None
            sheet.cell(row, 10).value = result.rank if result else None
        for column, title in enumerate([title for _, title in RESULT_COLUMNS] + ["Xếp loại"], start=6):
            if _blank(sheet.cell(1, column).value):
                sheet.cell(1, column).value = title
        with tempfile.NamedTemporaryFile(dir=source.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        workbook.close()
        os.replace(temporary, destination)
        return str(destination)
    finally:
        workbook.close()
        if temporary is not None:
            temporary.unlink(missing_ok=True)


RANK_CODES = {"Giỏi": "GIOI", "Khá": "KHA", "Đạt": "DAT", "Không đạt": "KDAT"}


class ScoreCard(QWidget):
    def __init__(self, person, pedestal_count, parent=None):
        super().__init__(parent)
        uic.loadUi(os.path.join(cf.DATA_DIR, "assets", "qt", "score_card.ui"), self)
        self.stt_label.setText(f"{person.number}.")
        self._name_text = f"{person.name} - {person.unit}"
        if pedestal_count == 1 and person.birthday:
            self._name_text += f" - {person.birthday}"
        self.name_label.setTextFormat(Qt.PlainText)
        self.name_label.setSizePolicy(QSizePolicy.Ignored, QSizePolicy.Fixed)
        self.name_label.setToolTip(self._name_text)
        self.name_label.setAccessibleName(self._name_text)
        self.name_label.installEventFilter(self)
        self._inactive_style = self.shoot_btn.styleSheet()
        self._active_style = self._inactive_style.replace("color: black;", "color: white;").replace(
            "background-color: rgb(255, 255, 255);", "background-color: rgb(0, 153, 0);").replace(
            "background-color: rgb(73, 220, 0);", "background-color: rgb(0, 184, 0);")
        self.shoot_btn.setToolTip(f"Vào bắn tại bệ {person.pedestal}")
        self.shoot_btn.setAccessibleName(f"{person.name}: vào bắn tại bệ {person.pedestal}")
        self.classify_label.setStyleSheet(self.classify_label.styleSheet() + "\npadding: 0;")
        self.classify_label.setFocusPolicy(Qt.NoFocus)
        self.update_result(0, None, False)

    def eventFilter(self, watched, event):
        if watched is self.name_label and event.type() == QEvent.Resize:
            self.name_label.setText(self.name_label.fontMetrics().elidedText(
                self._name_text, Qt.ElideRight, max(1, self.name_label.width())))
        return super().eventFilter(watched, event)

    def update_result(self, targets, rank, selected):
        self.targets_label.setText(
            f'<span style="color:#ff0000;">{targets}</span><span style="color:#000000;">/4</span>')
        self.targets_label.setToolTip(f"Bắn trúng {targets}/4 mục tiêu")
        self.classify_label.setText(RANK_CODES.get(rank, ""))
        self.classify_label.setToolTip(rank or "Chưa có xếp loại")
        self.shoot_btn.setStyleSheet(self._active_style if selected else self._inactive_style)


class ShootingRoster(QObject):
    loaded = pyqtSignal(object)
    saved = pyqtSignal(object)

    def __init__(self, window):
        super().__init__(window)
        self.window = window
        self.people = {}
        self.assignments = {}
        self.results = {}
        self.session_id = None
        self.session_assignments = {}
        self.loading = False
        self.saving = False
        self.closed = False
        self.file_path = ""
        self._pool = ThreadPoolExecutor(max_workers=1, thread_name_prefix="Roster")
        self.loaded.connect(self._finish_load)
        self.saved.connect(self._finish_save)
        window.select_file_btn.clicked.connect(self.choose_file)
        window.excel_file_option.triggered.connect(self.choose_file)
        window.new_route_btn.clicked.connect(self.next_group)
        window.new_route_btn.setToolTip("Xếp người chưa bắn tiếp theo vào mỗi bệ, theo thứ tự trong file")
        window.save_result_btn.clicked.connect(self.save_results)
        window.save_result_btn.setToolTip("Lưu file Kết quả - …xlsx cùng thư mục danh sách")
        for widget in window.client_widgets:
            widget.go_out_btn.clicked.disconnect(widget.reset_shooter)
            widget.go_out_btn.clicked.connect(lambda _checked=False, port=widget.id: self.clear(port))
            widget.go_out_btn.setToolTip("Bỏ người đang được chọn ở bệ này")
        self.cards = {}
        self._rebuild_list()

    def pending(self):
        window = self.window
        sid = window._active_review_session_id if window.testing else window._previous_review_session_id
        return any(count > len(window._review_sessions.get(sid, {}).get(port, []))
                   for (session, port), count in window._review_shot_counters.items() if session == sid)

    def locked(self):
        return self.loading or self.saving or self.window.testing or self.pending()

    def choose_file(self):
        if self.locked():
            return
        path, _ = QFileDialog.getOpenFileName(
            self.window, "Chọn danh sách bắn", self.file_path or cf.DATA_DIR, "Excel (*.xlsx)")
        if path:
            self.load_file(path)

    def load_file(self, path):
        if self.locked() or self.closed:
            return False
        self.loading = True
        self.refresh_controls()
        future = self._pool.submit(read_roster, path, self.window.p_num)

        def done(job):
            try:
                payload = (path, job.result(), None)
            except Exception as exc:
                payload = (path, None, str(exc))
            if not self.closed:
                try:
                    self.loaded.emit(payload)
                except RuntimeError:  # Parent was destroyed while the file was being read.
                    pass
        future.add_done_callback(done)
        return True

    def _finish_load(self, payload):
        if self.closed:
            return
        path, people, error = payload
        self.loading = False
        if error:
            self.refresh_controls()
            QMessageBox.warning(self.window, "Không nạp được danh sách",
                                f"{error}\n\nDùng file .xlsx có các cột: STT, Họ tên, Đơn vị, Ngày sinh.")
            return
        # Commit the new list only after the entire workbook has been validated.
        self.people = {person.number: person for person in people}
        self.file_path = path
        self.assignments.clear()
        self.results = {person.number: person.result for person in people if person.result is not None}
        self.session_id = None
        self.session_assignments.clear()
        for widget in self.window.client_widgets:
            self._show_person(widget.id, None)
        self._rebuild_list()

    def _rebuild_list(self):
        layout = self.window.score_card_layout
        while layout.count():
            item = layout.takeAt(0)
            if item.widget() is not None:
                item.widget().deleteLater()
        self.cards.clear()
        for person in self.people.values():
            card = ScoreCard(person, self.window.p_num)
            layout.addWidget(card)
            self.cards[person.number] = card
            card.shoot_btn.clicked.connect(lambda _checked=False, number=person.number: self.assign(number))
            self._refresh_item(person.number)
        layout.setAlignment(Qt.AlignTop)
        self.window.label_17.setText("Danh sách bắn")
        self.window.label_17.setToolTip(f"{len(self.people)} người\n{self.file_path}")
        self.window.select_file_btn.setToolTip(self.file_path or "Nạp STT, Họ tên, Đơn vị, Ngày sinh từ Excel")
        self.refresh_controls()

    def _refresh_item(self, number):
        person = self.people[number]
        selected = self.assignments.get(person.pedestal) == number
        result = self.results.get(number)
        targets = len(result.hit_targets) if result else 0
        if self.window.testing and self.session_assignments.get(person.pedestal) == number:
            targets = len(self._session_result(person.pedestal).hit_targets)
        self.cards[number].update_result(targets, result.rank if result else None, selected)

    def refresh_controls(self):
        enabled = not self.locked()
        window = self.window
        reason = ("Đang đọc danh sách Excel" if self.loading else "Đang lưu kết quả" if self.saving
                  else "Đang thực hiện bài bắn" if window.testing
                  else "Đang chấm nốt ảnh của bài vừa kết thúc" if self.pending() else "")
        window.select_file_btn.setEnabled(enabled)
        window.select_file_btn.setText("Đang đọc Excel…" if self.loading else "Chọn danh sách bắn")
        window.excel_file_option.setEnabled(enabled)
        window.new_route_btn.setEnabled(enabled and bool(self.people))
        window.new_route_btn.setToolTip(reason or "Xếp người chưa bắn tiếp theo vào mỗi bệ, theo thứ tự trong file")
        for card in self.cards.values():
            card.shoot_btn.setEnabled(enabled)
        window.save_result_btn.setEnabled(enabled and bool(self.file_path))
        window.save_result_btn.setText("Đang lưu…" if self.saving else "Lưu kết quả")
        for widget in window.client_widgets:
            widget.go_out_btn.setEnabled(enabled)
        # The start button remains usable to stop an active exercise.
        ready = not self.loading and not self.saving and not self.pending()
        window.start_btn.setEnabled(window.testing or ready)
        window.start_btn.setToolTip(reason if not window.testing and reason else "")

    def _show_person(self, port, number):
        widget = self.window.client_widgets[port - 1]
        self.window._result_sound.stop()
        widget.reset_results()
        widget.reset_shooter()
        if number is not None:
            person = self.people[number]
            widget.set_shooter(f"{person.name} - {person.unit}")
        widget.set_review_available(False)

    def assign(self, number):
        if self.locked() or number not in self.people:
            return False
        port = self.people[number].pedestal
        old = self.assignments.get(port)
        self.assignments[port] = number
        self._show_person(port, number)
        if old is not None:
            self._refresh_item(old)
        self._refresh_item(number)
        self.refresh_controls()
        return True

    def clear(self, port):
        if self.locked():
            return
        old = self.assignments.pop(port, None)
        self._show_person(port, None)
        if old is not None:
            self._refresh_item(old)
        self.refresh_controls()

    def next_group(self):
        if self.locked():
            return
        selected = {}
        for person in self.people.values():
            if person.number not in self.results:
                selected.setdefault(person.pedestal, person.number)
        for widget in self.window.client_widgets:
            self.clear(widget.id)
        for number in selected.values():
            self.assign(number)

    def begin_session(self, session_id):
        self.session_id = session_id
        self.session_assignments = dict(self.assignments)
        for number in self.session_assignments.values():
            self.results.pop(number, None)
            self._refresh_item(number)
        self.refresh_controls()

    def matches_session(self, session_id, port):
        return (not self.people or (session_id == self.session_id and
                self.assignments.get(port) == self.session_assignments.get(port)))

    def update_results(self, session_id):
        if session_id != self.session_id:
            return
        for port, number in self.session_assignments.items():
            if not self.window.testing:
                self.results[number] = self._session_result(port)
            self._refresh_item(number)
        self.refresh_controls()

    def _session_result(self, port):
        shots = self.window._review_sessions.get(self.session_id, {}).get(port, [])[:16]
        hits = {shot["target_index"] for shot in shots if shot.get("hit") is True}
        return RosterResult(frozenset(hits))

    def save_results(self):
        if self.locked() or self.closed or not self.file_path:
            return False
        self.saving = True
        self.refresh_controls()
        future = self._pool.submit(save_roster_results, self.file_path,
                                   tuple(self.people.values()), dict(self.results))

        def done(job):
            try:
                payload = (job.result(), None)
            except Exception as exc:
                payload = (None, str(exc))
            if not self.closed:
                try:
                    self.saved.emit(payload)
                except RuntimeError:
                    pass
        future.add_done_callback(done)
        return True

    def _finish_save(self, payload):
        if self.closed:
            return
        self.saving = False
        self.refresh_controls()
        path, error = payload
        if error:
            QMessageBox.warning(self.window, "Không lưu được kết quả",
                                f"{error}\n\nNếu file kết quả đang mở trong Excel, hãy đóng file rồi thử lại.")
        else:
            QMessageBox.information(self.window, "Lưu kết quả", f"Đã lưu kết quả:\n{path}")

    def shutdown(self):
        self.closed = True
        self._pool.shutdown(wait=False, cancel_futures=True)
